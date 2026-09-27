"""Documentul „Minutes of Meeting”: HTML (corpul emailului) + DOCX (atașament) + Markdown, în RO / RU / EN.

Conținutul urmează cerința Medpark: rezumatul ședinței, deciziile luate și sarcinile (action items) cu responsabil
și termen. Transcrierea nu intră în document. Formatare clinică (app/clinical.py): abrevieri standard și bold doar pe
valorile anormale și doze; la consiliul medical deciziile și sarcinile poartă pacientul (patul) la care se referă.
"""
from __future__ import annotations

import html
import re
from pathlib import Path

from rapidfuzz import fuzz

from .clinical import runs
from .config import get_config
from .llm import _ts_seconds, fmt_ts

L = {
    "ro": dict(doc="Proces-verbal", date="Data", duration="Durata", participants="Participanți",
               summary="Rezumat", decisions="Decizii", tasks="Sarcini", task="Sarcina", owner="Responsabil",
               deadline="Termen", not_set="nestabilit", urgent="urgent",
               no_decisions="Nu au fost luate decizii explicite.", no_tasks="Nu au fost stabilite sarcini.",
               general="Discuție generală", patient="Pacient", bed="Patul {n}", bed_long="Pacient — Patul {n}",
               footer="Generat automat de Secure MOM — procesare 100% locală, fără transmiterea datelor "
                      "în afara rețelei spitalului."),
    "ru": dict(doc="Протокол", date="Дата", duration="Длительность", participants="Участники",
               summary="Краткое содержание", decisions="Решения", tasks="Поручения", task="Поручение",
               owner="Ответственный", deadline="Срок", not_set="не установлен", urgent="срочно",
               no_decisions="Явных решений не принято.", no_tasks="Поручений не назначено.",
               general="Общее обсуждение", patient="Пациент", bed="Койка {n}", bed_long="Пациент — койка {n}",
               footer="Сгенерировано автоматически Secure MOM — 100% локальная обработка, данные не покидают "
                      "сеть больницы."),
    "en": dict(doc="Minutes of Meeting", date="Date", duration="Duration", participants="Participants",
               summary="Summary", decisions="Decisions", tasks="Action items", task="Task", owner="Owner",
               deadline="Deadline", not_set="not set", urgent="urgent",
               no_decisions="No explicit decisions were made.", no_tasks="No action items were assigned.",
               general="General discussion", patient="Patient", bed="Bed {n}", bed_long="Patient — Bed {n}",
               footer="Generated automatically by Secure MOM — 100% on-premise processing, no data leaves the "
                      "hospital network."),
}
# blocul de verificare umană — opțional (output.review_block în config), implicit în afara documentului
REVIEW = {
    "ro": dict(title="⚠️ Necesită verificare umană:", asr="recunoaștere audio nesigură",
               partial="decizie regăsită doar parțial în transcriere", listen="ascultați fragmentul"),
    "ru": dict(title="⚠️ Требует проверки человеком:", asr="неуверенное распознавание речи",
               partial="решение найдено в стенограмме лишь частично", listen="прослушайте фрагмент"),
    "en": dict(title="⚠️ Requires human verification:", asr="uncertain speech recognition",
               partial="decision only partially found in the transcript", listen="listen to the fragment"),
}


def _e(x) -> str:
    return html.escape(str(x)) if x not in (None, "") else "—"


def _date(iso: str | None, lang: str) -> str:
    """2026-09-27 -> 27.09.2026 (RO/RU); în engleză rămâne ISO."""
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", iso or "")
    return f"{m[3]}.{m[2]}.{m[1]}" if m and lang != "en" else (iso or "")


def _speaker_label(spk: str | None, mom: dict) -> str:
    if not spk:
        return ""
    for p in mom.get("participants", []):
        if p.get("speaker") == spk and p.get("name"):
            return f"{p['name']} ({spk})"
    return spk


def _owner(it: dict, mom: dict, lang: str) -> str:
    o, s = it.get("owner"), it.get("owner_speaker")
    if o and s:
        return f"{o} ({s})"
    return o or _speaker_label(s, mom) or L[lang]["not_set"]


def _deadline(it: dict, lang: str) -> str:
    d, t = _date(it.get("deadline_date"), lang), it.get("deadline_text")
    if d and t and t not in d:
        return f"{d} ({t})"
    return d or t or L[lang]["not_set"]


def meta(job: dict, lang: str) -> dict:
    cfg = get_config()
    return {
        "type": cfg.meeting_types.get(job["meeting_type"], {}).get(lang, job["meeting_type"]),
        "date": _date(job["meeting_date"], lang),
        "duration": fmt_ts(job.get("duration_s", 0)),
    }


def participants_line(mom: dict) -> str:
    out = []
    num = lambda p: int(re.sub(r"\D", "", p.get("speaker") or "") or 999)  # noqa: E731 — S1, S2, … în ordine
    for p in sorted(mom.get("participants", []), key=num):
        who = " · ".join(x for x in (p.get("name"), p.get("role")) if x)
        out.append(f"{p.get('speaker')}{' — ' + who if who else ''}")
    return "; ".join(out)


_SPEC = re.compile(r"\s*\((?:probabil|posibil|likely|possibly|вероятно|возможно|(?:în )?transcrierea?\b|"
                   r"in the transcript|в стенограмме)[^)]*(?:\)|$)", re.IGNORECASE)


def _clean(x: str | None) -> str:
    """Scoate speculațiile LLM-ului dintre paranteze („(probabil 8 mcg/kg/min)”) — documentul conține doar ce s-a spus."""
    return _SPEC.sub("", x or "").strip()


def _cap(x: str) -> str:
    return x[:1].upper() + x[1:] if x else x


def _labels(it: dict, lang: str) -> tuple[str, str, str, str]:
    """(cheie, tip, etichetă lungă „Pacient — Patul 8”, etichetă scurtă „Patul 8”) pentru subiectul unui element."""
    t = L[lang]
    k = it.get("subject_key") or "general"
    kind = it.get("subject_kind") or "theme"
    if k.startswith("bed:"):
        return k, kind, t["bed_long"].format(n=k[4:]), t["bed"].format(n=k[4:])
    if k == "general":
        return k, kind, t["general"], ""
    short = _cap((it.get("subject") or it.get("title") or k[4:]).strip())
    label = short if kind != "patient" or short.lower().startswith(t["patient"].lower()) else f"{t['patient']} — {short}"
    return k, kind, label, short


def _order(b: dict) -> tuple:
    """Pacienții întâi (în ordinea discuției), apoi temele, la final discuția generală."""
    return b["key"] == "general", b["kind"] != "patient", b["t"]


def subject_blocks(mom: dict, lang: str) -> list[dict]:
    """Subiectele ședinței în ordinea discuției — pacienții (după pat), apoi temele — cu deciziile și sarcinile lor."""
    blocks: dict[str, dict] = {}

    def block(it: dict) -> dict:
        k, kind, label, short = _labels(it, lang)
        b = blocks.get(k)
        if b is None:
            b = blocks[k] = {"key": k, "kind": kind, "label": label, "short": short, "t": 10 ** 9,
                             "decisions": [], "plan": []}
        b["t"] = min(b["t"], _ts_seconds(it.get("timestamp")) or 0)
        return b

    for tp in mom.get("topics", []):
        block(tp)
    plan_all = mom.get("action_items", [])
    for d in mom.get("decisions", []):
        # „Se solicită consult urolog” și sarcina „Consult urolog” din același moment = același lucru: rămâne sarcina
        if any((a.get("timestamp") == d.get("timestamp") and
                fuzz.token_set_ratio(a.get("task") or "", d.get("decision") or "") >= 35) or
               fuzz.token_set_ratio(a.get("task") or "", d.get("decision") or "") >= 80 for a in plan_all):
            continue
        block(d)["decisions"].append(d)
    for a in plan_all:
        block(a)["plan"].append(a)
    out = [b for b in blocks.values() if b["decisions"] or b["plan"]]
    return sorted(out, key=_order)


def summary_items(mom: dict, lang: str) -> list[tuple[str, str]]:
    """Rezumatul: când s-au discutat pacienți, starea fiecăruia (parametri, evoluție, tratament), grupată după pat;
    altfel rezumatul general al ședinței. Lista de (etichetă, text); eticheta e goală pentru rezumatul general."""
    topics = [tp for tp in mom.get("topics", []) if (tp.get("summary") or "").strip()]
    if not any(tp.get("subject_kind") == "patient" for tp in topics):
        s = mom.get("summary") or " ".join(tp.get("summary") or "" for tp in topics)
        return [("", _clean(s))] if s.strip() else []
    # LLM-ul mai scrie pentru același pacient și „situații” repetate (sarcini, confuzii): din fiecare bucată a
    # transcrierii rămâne doar cea mai completă situație clinică a pacientului
    best: dict[tuple, dict] = {}
    for tp in topics:
        k = tp.get("subject_key") or "general"
        cur = best.get((k, tp.get("part")))
        if cur is None or len(tp.get("summary") or "") > len(cur.get("summary") or ""):
            best[(k, tp.get("part"))] = tp
    groups: dict[str, dict] = {}
    for tp in sorted(best.values(), key=lambda x: _ts_seconds(x.get("timestamp")) or 0):
        k, kind, _, short = _labels(tp, lang)
        g = groups.setdefault(k, {"key": k, "kind": kind, "short": short, "t": 10 ** 9, "texts": []})
        g["t"] = min(g["t"], _ts_seconds(tp.get("timestamp")) or 0)
        g["texts"].append(_clean(tp.get("summary")))
    return [(g["short"] or L[lang]["general"], " ".join(g["texts"])) for g in sorted(groups.values(), key=_order)]


def decision_groups(mom: dict, lang: str) -> list[tuple[str | None, list[dict]]]:
    """Deciziile pe subiecte; fără subtitlu când ședința are un singur subiect."""
    groups = [(b["label"], b["decisions"]) for b in subject_blocks(mom, lang) if b["decisions"]]
    if len(groups) == 1 and not any(b["kind"] == "patient" for b in subject_blocks(mom, lang)):
        return [(None, groups[0][1])]
    return groups


def task_rows(mom: dict, lang: str) -> list[dict]:
    blocks = subject_blocks(mom, lang)
    many = len(blocks) > 1 or any(b["kind"] == "patient" for b in blocks)
    return [{"subject": b["short"] if many else "", "task": a.get("task"), "owner": _owner(a, mom, lang),
             "deadline": _deadline(a, lang), "urgent": a.get("priority") == "urgent"}
            for b in blocks for a in b["plan"]]


def review_items(mom: dict, lang: str) -> list[dict]:
    """Confuzii din discuție, cuvinte critice nesigure, surse parțiale — pentru blocul opțional de verificare."""
    c = REVIEW[lang]
    labels = {b["key"]: b["label"] for b in subject_blocks(mom, lang)}
    out, seen = [], set()
    for n in mom.get("verification_notes", []) or []:
        out.append({"t": n.get("timestamp"), "subject": labels.get(n.get("subject_key"), ""), "text": _clean(n.get("note"))})
    for key, field in (("decisions", "decision"), ("action_items", "task")):
        for it in mom.get(key, []):
            subj = labels.get(it.get("subject_key"), "")
            u = it.get("uncertain_terms") or []
            words = ", ".join(f"„{x['word']}” ({round(x['p'] * 100)}%)" for x in u)
            if u and (it.get("timestamp"), words) not in seen:
                seen.add((it.get("timestamp"), words))
                out.append({"t": it.get("timestamp"), "subject": subj, "text": f"{c['asr']}: {words} — {c['listen']}"})
            if not it.get("verified", True):
                out.append({"t": it.get("timestamp"), "subject": subj, "text": f"{c['partial']}: {_clean(it.get(field))}"})
    return sorted(out, key=lambda x: _ts_seconds(x.get("t")) or 0)


def _review_on() -> bool:
    return bool(get_config().output.get("review_block", False))


def _rich_html(text: str | None, lang: str) -> str:
    return "".join(f"<b>{html.escape(x)}</b>" if b else html.escape(x) for x, b in runs(_clean(text), lang))


CSS = """
body{margin:0;padding:28px 14px;background:#eef2f7;color:#18212f;
     font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,Helvetica,Arial,sans-serif}
.doc{max-width:820px;margin:0 auto;background:#fff;border-radius:14px;padding:34px 40px 24px;
     box-shadow:0 1px 2px rgba(16,34,64,.06),0 10px 30px rgba(16,34,64,.07)}
.kicker{font-size:11px;letter-spacing:.14em;text-transform:uppercase;color:#0e7c86;font-weight:700}
h1{font-size:23px;line-height:1.28;margin:6px 0 12px;color:#0d1b2e}
.meta{font-size:13px;color:#5a6577;line-height:1.7;border-bottom:1px solid #e6eaf0;padding-bottom:14px}
.meta b{color:#1f2a3a;font-weight:600}
h2{font-size:12px;letter-spacing:.12em;text-transform:uppercase;color:#0b5394;margin:26px 0 8px}
h3{font-size:14.5px;margin:14px 0 2px;color:#1f2a3a}
p,li,td{font-size:14.5px;line-height:1.55} p{margin:4px 0} p.sum{margin:0 0 10px}
ol{margin:4px 0 0;padding-left:24px} li{margin:5px 0;padding-left:2px}
table{width:100%;border-collapse:collapse;margin-top:2px}
th{font-size:10.5px;text-transform:uppercase;letter-spacing:.08em;color:#6b7587;text-align:left;font-weight:700;
   padding:6px 8px;border-bottom:2px solid #e3e8ef}
td{padding:10px 8px;border-bottom:1px solid #edf0f4;vertical-align:top}
td.box{width:16px;padding-right:0;color:#7b8597;font-size:16px;line-height:1.3}
td.who{color:#2a3546;min-width:90px} td.due{color:#2a3546;white-space:nowrap}
.na{color:#9aa3b2;font-style:italic}
.subj{display:inline-block;font-size:11.5px;font-weight:600;color:#0b5394;background:#eaf2fb;border-radius:6px;
      padding:1px 7px;margin:0 6px 2px 0}
.urg{display:inline-block;font-size:10px;font-weight:700;letter-spacing:.06em;color:#b42318;background:#fdecea;
     border-radius:6px;padding:1px 6px;margin-left:6px;text-transform:uppercase}
.none{color:#8a93a3;font-size:14px}
.review{margin:18px 0 0;padding:12px 16px;border-radius:10px;background:#fff7e8;border:1px solid #f3cf86}
.review b.t{display:block;color:#8a4b00;margin-bottom:4px} .review ul{margin:0;padding-left:18px}
.review li{font-size:13px;margin:3px 0}
.foot{margin-top:28px;font-size:11px;color:#8a93a3}
@media (max-width:600px){.doc{padding:24px 18px}td.due{white-space:normal}}
"""


def render_html(mom: dict, job: dict, lang: str) -> str:
    t = L[lang]
    m = meta(job, lang)
    na = t["not_set"]
    h = [f"<!doctype html><html lang='{lang}'><head><meta charset='utf-8'>"
         f"<meta name='viewport' content='width=device-width,initial-scale=1'><title>{_e(mom.get('title'))}</title>"
         f"<style>{CSS}</style></head><body><div class='doc'>",
         f"<div class='kicker'>{html.escape(t['doc'])} · {html.escape(m['type'])}</div>",
         f"<h1>{_e(mom.get('title') or t['doc'])}</h1>",
         f"<div class='meta'>{t['date']}: <b>{_e(m['date'])}</b> &nbsp;·&nbsp; {t['duration']}: <b>{m['duration']}</b>"]
    pl = participants_line(mom)
    if pl:
        h.append(f"<br>{t['participants']}: <b>{html.escape(pl)}</b>")
    h.append("</div>")

    if _review_on():
        rv = review_items(mom, lang)
        if rv:
            h.append(f"<div class='review'><b class='t'>{REVIEW[lang]['title']}</b><ul>" + "".join(
                f"<li>[{_e(x.get('t'))}] {('<b>' + html.escape(x['subject']) + '</b>: ') if x.get('subject') else ''}"
                f"{_e(x['text'])}</li>" for x in rv) + "</ul></div>")

    items = summary_items(mom, lang)
    if items:
        h.append(f"<h2>{t['summary']}</h2>")
        for label, text in items:
            tag = f"<span class='subj'>{html.escape(label)}</span>" if label else ""
            h.append(f"<p class='sum'>{tag}{_rich_html(text, lang)}</p>")

    h.append(f"<h2>{t['decisions']}</h2>")
    groups = decision_groups(mom, lang)
    n = 1
    for label, items in groups:
        if label:
            h.append(f"<h3>{html.escape(label)}</h3>")
        h.append(f"<ol start='{n}'>" + "".join(f"<li>{_rich_html(d.get('decision'), lang)}</li>" for d in items) + "</ol>")
        n += len(items)
    if not groups:
        h.append(f"<p class='none'>{t['no_decisions']}</p>")

    h.append(f"<h2>{t['tasks']}</h2>")
    rows = task_rows(mom, lang)
    if rows:
        h.append(f"<table><tr><th></th><th>{t['task']}</th><th>{t['owner']}</th><th>{t['deadline']}</th></tr>")
        for r in rows:
            subj = f"<span class='subj'>{html.escape(r['subject'])}</span>" if r["subject"] else ""
            urg = f"<span class='urg'>{t['urgent']}</span>" if r["urgent"] else ""
            cell = lambda v: f"<span class='na'>{na}</span>" if v == na else html.escape(v)  # noqa: E731
            h.append(f"<tr><td class='box'>☐</td><td>{subj}{_rich_html(r['task'], lang)}{urg}</td>"
                     f"<td class='who'>{cell(r['owner'])}</td><td class='due'>{cell(r['deadline'])}</td></tr>")
        h.append("</table>")
    else:
        h.append(f"<p class='none'>{t['no_tasks']}</p>")

    h.append(f"<div class='foot'>{t['footer']}</div></div></body></html>")
    return "".join(h)


def render_docx(mom: dict, job: dict, lang: str, path: Path) -> Path:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    t = L[lang]
    m = meta(job, lang)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    teal, blue, grey, ink = (RGBColor(0x0E, 0x7C, 0x86), RGBColor(0x0B, 0x53, 0x94), RGBColor(0x6B, 0x75, 0x87),
                             RGBColor(0x1F, 0x2A, 0x3A))

    def para(text: str = "", size: float | None = None, color=None, bold: bool = False, style: str | None = None):
        p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
        if text:
            r = p.add_run(text)
            r.bold = bold
            if size:
                r.font.size = Pt(size)
            if color is not None:
                r.font.color.rgb = color
        return p

    def rich(p, text: str | None, size: float | None = None):
        for x, b in runs(_clean(text), lang):
            r = p.add_run(x)
            r.bold = b
            if size:
                r.font.size = Pt(size)

    def section(title: str):
        p = para(title.upper(), size=10, color=blue, bold=True)
        p.paragraph_format.space_before = Pt(14)

    para(f"{t['doc']} · {m['type']}".upper(), size=8.5, color=teal, bold=True)
    doc.add_heading(mom.get("title") or t["doc"], level=0)
    para(f"{t['date']}: {m['date']}  ·  {t['duration']}: {m['duration']}", size=9.5, color=grey)
    pl = participants_line(mom)
    if pl:
        para(f"{t['participants']}: {pl}", size=9.5, color=grey)

    if _review_on():
        rv = review_items(mom, lang)
        if rv:
            para(REVIEW[lang]["title"], size=11, color=RGBColor(0x8A, 0x4B, 0x00), bold=True)
            for x in rv:
                p = para(style="List Bullet")
                p.add_run(f"[{x.get('t')}] ").font.color.rgb = grey
                if x.get("subject"):
                    p.add_run(f"{x['subject']}: ").bold = True
                p.add_run(x["text"])

    items = summary_items(mom, lang)
    if items:
        section(t["summary"])
        for label, text in items:
            p = doc.add_paragraph()
            if label:
                r = p.add_run(f"{label} · ")
                r.bold, r.font.color.rgb = True, blue
            rich(p, text)

    section(t["decisions"])
    groups = decision_groups(mom, lang)
    for label, items in groups:
        if label:
            para(label, size=11, color=ink, bold=True)
        for d in items:
            rich(para(style="List Number"), d.get("decision"))
    if not groups:
        para(t["no_decisions"], color=grey)

    section(t["tasks"])
    rows = task_rows(mom, lang)
    if rows:
        tb = doc.add_table(rows=1, cols=4)
        tb.style = "Table Grid"
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        widths = [Cm(0.8), Cm(9.6), Cm(3.4), Cm(3.2)]
        for i, (cell, name) in enumerate(zip(tb.rows[0].cells, ["", t["task"], t["owner"], t["deadline"]])):
            cell.width = widths[i]
            r = cell.paragraphs[0].add_run(name)
            r.bold, r.font.size, r.font.color.rgb = True, Pt(9), grey
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "EEF3F9")
            cell._tc.get_or_add_tcPr().append(shd)
        for row in rows:
            cells = tb.add_row().cells
            for i, c in enumerate(cells):
                c.width = widths[i]
            cells[0].paragraphs[0].add_run("☐").font.size = Pt(12)
            p = cells[1].paragraphs[0]
            if row["subject"]:
                r = p.add_run(f"{row['subject']} · ")
                r.bold, r.font.color.rgb = True, blue
            rich(p, row["task"])
            if row["urgent"]:
                r = p.add_run(f"  {t['urgent'].upper()}")
                r.bold, r.font.size, r.font.color.rgb = True, Pt(8), RGBColor(0xB4, 0x23, 0x18)
            for i, key in ((2, "owner"), (3, "deadline")):
                r = cells[i].paragraphs[0].add_run(row[key])
                if row[key] == t["not_set"]:
                    r.italic, r.font.color.rgb = True, grey
    else:
        para(t["no_tasks"], color=grey)

    p = para(t["footer"], size=8, color=grey)
    p.paragraph_format.space_before = Pt(18)
    doc.save(path)
    return path


def render_markdown(mom: dict, job: dict, lang: str) -> str:
    t = L[lang]
    m = meta(job, lang)
    md = lambda x: "".join(f"**{a}**" if b else a for a, b in runs(_clean(x), lang))  # noqa: E731
    out = [f"*{t['doc']} · {m['type']}*", "", f"# {mom.get('title') or t['doc']}", "",
           f"{t['date']}: {m['date']} · {t['duration']}: {m['duration']}"]
    pl = participants_line(mom)
    if pl:
        out += ["", f"{t['participants']}: {pl}"]
    if _review_on():
        rv = review_items(mom, lang)
        if rv:
            out += ["", f"### {REVIEW[lang]['title']}"] + [
                f"- [{x.get('t')}] {x['subject'] + ': ' if x.get('subject') else ''}{x['text']}" for x in rv]
    items = summary_items(mom, lang)
    if items:
        out += ["", f"## {t['summary']}"] + [f"- **{label}** · {md(text)}" if label else md(text) for label, text in items]
    out += ["", f"## {t['decisions']}"]
    groups = decision_groups(mom, lang)
    n = 1
    for label, items in groups:
        if label:
            out += ["", f"### {label}"]
        for d in items:
            out.append(f"{n}. {md(d.get('decision'))}")
            n += 1
    if not groups:
        out.append(t["no_decisions"])
    out += ["", f"## {t['tasks']}"]
    rows = task_rows(mom, lang)
    if rows:
        out += [f"| | {t['task']} | {t['owner']} | {t['deadline']} |", "|---|---|---|---|"]
        for r in rows:
            task = (f"**{r['subject']}** · " if r["subject"] else "") + md(r["task"]).replace("|", "/") + \
                   (f" **{t['urgent'].upper()}**" if r["urgent"] else "")
            out.append(f"| ☐ | {task} | {r['owner']} | {r['deadline']} |")
    else:
        out.append(t["no_tasks"])
    out += ["", f"_{t['footer']}_"]
    return "\n".join(out)
